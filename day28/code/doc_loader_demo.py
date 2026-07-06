# -*- coding: utf-8 -*-
"""
Day 28 · Document Loader 演示

支持格式：PDF / Word / Markdown / Web(本地HTML) / CSV / TXT
运行：cd day28/code && python3 doc_loader_demo.py
"""

from __future__ import annotations

import os
import sys
from pathlib import Path
from typing import Any

CODE_DIR = Path(__file__).resolve().parent
SAMPLE_DIR = CODE_DIR / "data" / "sample_docs"
sys.path.insert(0, str(CODE_DIR))

from ensure_samples import ensure_all_samples  # noqa: E402

try:
    from dotenv import load_dotenv
except ImportError:  # pragma: no cover
    load_dotenv = None  # type: ignore[assignment]

from langchain_core.documents import Document


def load_dotenv_file() -> None:
    if load_dotenv is None:
        return
    for candidate in (CODE_DIR / ".env", CODE_DIR.parent / ".env"):
        if candidate.is_file():
            load_dotenv(candidate)
            return
    load_dotenv()


def summarize_docs(label: str, docs: list[Document]) -> None:
    print(f"\n{'=' * 60}")
    print(f"[{label}] 共 {len(docs)} 个 Document")
    print("=" * 60)
    for i, doc in enumerate(docs[:3]):
        preview = doc.page_content.replace("\n", " ")[:120]
        meta = doc.metadata
        print(f"  #{i} chars={len(doc.page_content)} source={meta.get('source', 'N/A')}")
        print(f"      preview: {preview}...")
    if len(docs) > 3:
        print(f"  ... 另有 {len(docs) - 3} 个 Document 未展示")


def load_markdown(path: Path) -> list[Document]:
    from langchain_community.document_loaders import TextLoader

    loader = TextLoader(str(path), encoding="utf-8")
    docs = loader.load()
    for d in docs:
        d.metadata["format"] = "markdown"
    return docs


def load_txt(path: Path) -> list[Document]:
    from langchain_community.document_loaders import TextLoader

    loader = TextLoader(str(path), encoding="utf-8")
    docs = loader.load()
    for d in docs:
        d.metadata["format"] = "txt"
    return docs


def load_pdf(path: Path) -> list[Document]:
    from langchain_community.document_loaders import PyPDFLoader

    loader = PyPDFLoader(str(path))
    docs = loader.load()
    total_chars = sum(len(d.page_content) for d in docs)
    if total_chars < 50:
        fallback = path.with_suffix(".pdf.txt")
        if fallback.is_file():
            print(f"[INFO] PDF 文本过少，回退 {fallback.name}")
            docs = load_txt(fallback)
    for d in docs:
        d.metadata["format"] = "pdf"
    return docs


def load_word(path: Path) -> list[Document]:
    try:
        from docx import Document as DocxDocument

        docx = DocxDocument(str(path))
        text = "\n".join(p.text for p in docx.paragraphs if p.text.strip())
        docs = [Document(page_content=text, metadata={"source": str(path), "format": "docx"})]
    except Exception:
        from langchain_community.document_loaders import Docx2txtLoader

        loader = Docx2txtLoader(str(path))
        docs = loader.load()
        for d in docs:
            d.metadata["format"] = "docx"
    return docs


def load_csv(path: Path) -> list[Document]:
    from langchain_community.document_loaders import CSVLoader

    loader = CSVLoader(
        file_path=str(path),
        encoding="utf-8",
        csv_args={"delimiter": ","},
    )
    docs = loader.load()
    for d in docs:
        d.metadata["format"] = "csv"
    return docs


def load_web_local(html_path: Path) -> list[Document]:
    """优先本地 HTML，避免 CI 无外网。"""
    from langchain_community.document_loaders import BSHTMLLoader

    loader = BSHTMLLoader(str(html_path), open_encoding="utf-8")
    docs = loader.load()
    for d in docs:
        d.metadata["format"] = "html"
    return docs


def load_web_remote(url: str) -> list[Document]:
    from langchain_community.document_loaders import WebBaseLoader

    loader = WebBaseLoader(url)
    return loader.load()


def run_all_loaders() -> dict[str, Any]:
    load_dotenv_file()
    ensure_all_samples()

    results: dict[str, Any] = {}

    md_path = SAMPLE_DIR / "product_guide.md"
    txt_path = SAMPLE_DIR / "faq_refund.txt"
    pdf_path = SAMPLE_DIR / "ebook.pdf"
    docx_path = SAMPLE_DIR / "policy.docx"
    csv_path = SAMPLE_DIR / "orders_sample.csv"
    html_path = SAMPLE_DIR / "sparktech_web.html"

    loaders_plan = [
        ("Markdown", lambda: load_markdown(md_path)),
        ("TXT", lambda: load_txt(txt_path)),
        ("PDF", lambda: load_pdf(pdf_path)),
        ("Word", lambda: load_word(docx_path)),
        ("CSV", lambda: load_csv(csv_path)),
        ("Web(本地HTML)", lambda: load_web_local(html_path)),
    ]

    all_docs: list[Document] = []
    for label, fn in loaders_plan:
        try:
            docs = fn()
            summarize_docs(label, docs)
            results[label] = {"count": len(docs), "ok": True}
            all_docs.extend(docs)
        except Exception as exc:  # pragma: no cover
            print(f"[FAIL] {label}: {exc}")
            results[label] = {"count": 0, "ok": False, "error": str(exc)}

  # 可选：外网 WebBaseLoader（默认跳过）
    if os.getenv("WEB_DEMO_URL") and os.getenv("SPARKTECH_MOCK", "1") != "1":
        try:
            remote_docs = load_web_remote(os.environ["WEB_DEMO_URL"])
            summarize_docs("Web(远程)", remote_docs)
            results["Web(远程)"] = {"count": len(remote_docs), "ok": True}
            all_docs.extend(remote_docs)
        except Exception as exc:
            print(f"[SKIP] Web(远程): {exc}")

    print(f"\n{'=' * 60}")
    print(f"合计加载 {len(all_docs)} 个 Document（来自 {len(results)} 种 Loader）")
    print("=" * 60)
    return results


def main() -> int:
    print("Day 28 · doc_loader_demo.py")
    print(f"样本目录: {SAMPLE_DIR}")
    results = run_all_loaders()
    failed = [k for k, v in results.items() if not v.get("ok")]
    if failed:
        print(f"\n[WARN] 部分 Loader 失败: {failed}")
        return 1
    print("\n[OK] 全部 Document Loader 演示完成")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
